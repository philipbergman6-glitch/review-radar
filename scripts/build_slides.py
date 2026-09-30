"""Build docs/slides/review-radar.pptx from the running order in docs/SLIDES.md.

docs/SLIDES.md stays the source of truth for the talk -- its timings and binding wording are
what tests/test_written_deliverables.py checks. This script is the deck the brief asks for:
a title, the eight slides in that running order, and one divider where the live demo sits.
Slide 2 is variant B, as called at the rehearsal of 2026-09-14. Every figure here is copied
from SLIDES.md / DESIGN.md; nothing is computed, so a number that changes there must be
changed here by hand.

Run through `make slides` (python-pptx is pulled in with `uv run --with`, not locked).
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "docs" / "slides" / "review-radar.pptx"

FONT = "Arial"
MONO = "Courier New"
PAPER = RGBColor(0xFA, 0xFA, 0xF7)
INK = RGBColor(0x14, 0x21, 0x3D)
MUTED = RGBColor(0x5B, 0x64, 0x72)
ACCENT = RGBColor(0xC8, 0x55, 0x3D)
PANEL = RGBColor(0xEE, 0xEC, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
MARGIN = 0.7


def text(slide, x, y, w, h, runs, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.08, gap=0):
    """One text box. `runs` is a string, or a list of paragraphs; a paragraph is a string or
    a list of (text, overrides) runs where overrides may set bold / color / size / font."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    paragraphs = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        for chunk in ([(para, {})] if isinstance(para, str) else para):
            body, over = chunk
            r = p.add_run()
            r.text = body
            r.font.name = over.get("font", font)
            r.font.size = Pt(over.get("size", size))
            r.font.bold = over.get("bold", bold)
            r.font.color.rgb = over.get("color", color)
    return box


def panel(slide, x, y, w, h, fill=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def card(slide, x, y, w, h, figure, caption, *, figure_size=30, accent=False):
    """A figure and the sentence that says what it measures."""
    panel(slide, x, y, w, h)
    panel(slide, x, y, 0.07, h, fill=ACCENT if accent else INK)
    text(slide, x + 0.25, y + 0.18, w - 0.4, 0.7, figure, size=figure_size, bold=True,
         color=ACCENT if accent else INK)
    text(slide, x + 0.25, y + 0.18 + figure_size / 72 * 1.5, w - 0.4, h - 1.0, caption,
         size=15, color=MUTED)


def node(slide, x, y, w, h, label, *, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = INK
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Inches(0.04))
    for i, line in enumerate(label.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(13 if i == 0 else 10.5)
        r.font.bold = i == 0
        r.font.color.rgb = INK if i == 0 else MUTED
    return box


def arrow(slide, x1, y1, x2, y2, *, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = INK
    line.line.width = Pt(1.5)
    ln = line.line._get_or_add_ln()
    if dashed:
        dash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(dash)
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return line


def base(prs, number, title, seconds, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    if title:
        text(slide, MARGIN, 0.5, W - 2 * MARGIN, 0.9, title, size=32, bold=True)
        panel(slide, MARGIN, 1.32, 0.9, 0.06, fill=ACCENT)
    if number is not None:
        text(slide, W - MARGIN - 3, H - 0.55, 3, 0.3, f"Review Radar · {number}", size=10,
             color=MUTED, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = f"[{seconds} s] {notes}"
    return slide


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    content_w = W - 2 * MARGIN

    # 0 -- title -------------------------------------------------------------------------
    s = base(prs, None, None, 10, "Name, course, the dataset in one breath. Move on.")
    panel(s, 0, 0, 0.35, H, fill=ACCENT)
    text(s, 1.2, 2.2, 11, 1.2, "Review Radar", size=66, bold=True)
    text(s, 1.2, 3.45, 11, 0.6, "Streaming review intelligence: which products are declining, and what are people complaining about?",
         size=22, color=MUTED)
    text(s, 1.2, 5.3, 11, 1.0,
         ["Amazon Reviews 2023 · All_Beauty · 701,528 reviews",
          "Philip Bergman · BIU 8688697201, Big Data and AI · final project (solo)"],
         size=16, color=INK, spacing=1.3)

    # 1 -- the question ------------------------------------------------------------------
    s = base(prs, 1, "The question, and who asks it", 30,
             "Read the question verbatim. It was fixed before any analysis ran (ADR-0001). "
             "A category manager owns a shelf and wants a shortlist, not a model.")
    panel(s, MARGIN, 1.85, 0.08, 1.9, fill=ACCENT)
    text(s, MARGIN + 0.4, 1.8, content_w - 0.4, 2.0,
         "Which products experienced a sustained decline in customer ratings, and which "
         "complaint themes increased during that decline?", size=32, bold=True, spacing=1.15)
    text(s, MARGIN + 0.4, 3.95, content_w - 0.4, 0.5,
         "Asked by a category manager. Fixed before the analysis was run.", size=18, color=MUTED)
    for i, (figure, caption) in enumerate([
        ("701,528", "reviews"), ("112,590", "catalogue products"), ("2000 → 2023", "2000-11-01 to 2023-09-09"),
    ]):
        card(s, MARGIN + i * 4.05, 4.9, 3.8, 1.6, figure, caption)

    # 2 -- the answer, variant B ---------------------------------------------------------
    s = base(prs, 2, "The answer, in one line", 40,
             "Variant B, the narrow claim. The fork on candidate supply did not fire (708 matched "
             "candidates, rule is < 3); the guard on measured quality did (THEMES_QUALITY macro-F1 "
             "0.4583 vs 0.70). Close with: '-- and here is why that narrowness is the result.'")
    text(s, MARGIN, 1.7, content_w, 1.3,
         "The completed result identifies sustained rating declines; aspect-level explanation is preliminary.",
         size=27, bold=True, spacing=1.15)
    card(s, MARGIN, 3.3, 3.8, 1.75, "139 of 502", "eligible products raised an alert during 2020–2023, under a rule frozen beforehand", accent=True)
    card(s, MARGIN + 4.05, 3.3, 3.8, 1.75, "0.80 / month", "false alerts on placebo data, against a budget of one")
    card(s, MARGIN + 8.1, 3.3, 3.83, 1.75, "18.6%", "of declines its own size are detected, so this is a floor, not a census")
    text(s, MARGIN, 5.4, content_w, 1.3,
         [[("Hero: ", {"bold": True}), ("B00RPJZMUM", {"font": MONO, "size": 15}),
           (" — hard_to_use and not_as_described give way to does_not_work (“motor stopped working”).", {})],
          [(("An illustration over 18 baseline and 3 recent labelled mentions, not an estimate. "
             "The decline claim stands; the theme claim is illustrative."), {"color": MUTED})]],
         size=17, gap=6)

    # 3 -- the data and the V's ----------------------------------------------------------
    s = base(prs, 3, "The data, and which V’s it exercises", 30,
             "Four V's, each with a measured figure. Velocity is a range on purpose: the host is "
             "shared and a point estimate would not reproduce.")
    cards = [
        ("Volume", "311 MB → 96.5 MB", "bronze Parquet, 3.2× smaller · 0.54 GB raw JSONL · 693,547 ES docs · 345,418 vectors"),
        ("Velocity", "232k–484k rec/s", "producer throughput, a range · the stream replay is paced to a frozen 3,000 rec/s, measured 2,999"),
        ("Variety", "free text + free-form JSON", "a details dict whose keys collide case-insensitively · price as null, “9.99”, “$9.99” or a range · Postgres catalogue joined over JDBC"),
        ("Veracity", "6,139 collision groups", "over 13,415 rows: 6,138 byte-exact, 1 conflicting on helpful_vote, 0 disagreeing on rating → 7,276 rows removed"),
    ]
    for i, (label, figure, caption) in enumerate(cards):
        x = MARGIN + (i % 2) * 6.08
        y = 1.75 + (i // 2) * 2.55
        panel(s, x, y, 5.85, 2.3)
        panel(s, x, y, 0.07, 2.3, fill=INK)
        text(s, x + 0.25, y + 0.15, 5.4, 0.3, label.upper(), size=11, bold=True, color=ACCENT)
        text(s, x + 0.25, y + 0.48, 5.4, 0.6, figure, size=26, bold=True)
        text(s, x + 0.25, y + 1.1, 5.4, 1.15, caption, size=14.5, color=MUTED)

    # 4 -- architecture ------------------------------------------------------------------
    s = base(prs, 4, "Architecture and data flow", 40,
             "Walk left to right. Postgres does three jobs: Iceberg catalogue (the atomic pointer "
             "swap an object store cannot do), JDBC enrichment source, and the run ledger. Then the "
             "exactly-once proof as a figure: make eos SIGKILLs bronze mid-stream and restarts from "
             "the checkpoint.")
    y0, h = 2.75, 0.75
    node(s, 0.55, y0, 1.55, h, "raw JSONL\n701,528 reviews", shape=MSO_SHAPE.FLOWCHART_MAGNETIC_DISK)
    node(s, 2.5, y0, 1.6, h, "replay producer")
    node(s, 4.5, y0, 1.35, h, "Kafka", fill=PANEL)
    node(s, 6.3, y0, 2.2, h, "Spark batch\nbronze → silver → gold")
    node(s, 9.0, y0, 2.0, h, "Elasticsearch\nreviews · product_month", fill=PANEL)
    node(s, 11.45, y0, 1.3, h, "Kibana")
    node(s, 6.3, 1.6, 2.2, h, "Spark streaming\nwatermark 30 d")
    node(s, 9.0, 1.6, 2.0, h, "Iceberg on MinIO", fill=PANEL, shape=MSO_SHAPE.FLOWCHART_MAGNETIC_DISK)
    node(s, 3.6, 3.95, 2.45, h, "PostgreSQL\ncatalogue · products · run ledger", shape=MSO_SHAPE.FLOWCHART_MAGNETIC_DISK)
    node(s, 8.2, 3.95, 3.6, h, "AI layer\nkNN · hybrid RRF · theme labelling · RAG", fill=WHITE)
    mid = y0 + h / 2
    arrow(s, 2.1, mid, 2.5, mid)
    arrow(s, 4.1, mid, 4.5, mid)
    arrow(s, 5.85, mid, 6.3, mid)
    arrow(s, 8.5, mid, 9.0, mid)
    arrow(s, 11.0, mid, 11.45, mid)
    arrow(s, 5.5, y0, 6.3, 1.6 + h / 2)            # Kafka -> streaming
    arrow(s, 8.5, 1.6 + h / 2, 9.0, 1.6 + h / 2)   # streaming -> Iceberg
    arrow(s, 8.3, y0, 9.2, 1.6 + h)                # batch -> Iceberg
    arrow(s, 6.05, 4.15, 6.9, y0 + h, dashed=True) # Postgres -> batch
    text(s, 5.75, 3.55, 0.8, 0.25, "JDBC", size=11, color=MUTED)
    arrow(s, 10.0, y0 + h, 10.0, 3.95)             # ES -> AI
    text(s, MARGIN, 5.05, content_w, 0.3, "EXACTLY-ONCE, AS A FIGURE", size=11, bold=True, color=ACCENT)
    for i, (figure, caption) in enumerate([
        ("120,000", "records expected"), ("89,964", "rows before the SIGKILL"),
        ("120,000", "rows after the restart"), ("0", "duplicate (partition, offset) pairs"),
    ]):
        x = MARGIN + i * 3.04
        panel(s, x, 5.4, 2.85, 1.25)
        text(s, x + 0.2, 5.47, 2.5, 0.6, figure, size=26, bold=True, color=ACCENT if i == 3 else INK)
        text(s, x + 0.2, 6.1, 2.5, 0.5, caption, size=14, color=MUTED)

    # 5 -- the AI capability -------------------------------------------------------------
    s = base(prs, 5, "The AI capability, and why this one", 35,
             "The course names the gap itself in deck 5. We measure it in the same engine. "
             "H-E1 is NOT held: kNN alone does not beat BM25; the hybrid is where the gain is. "
             "Themes and RAG sit above the same retriever, run by a local qwen3:8b.")
    panel(s, MARGIN, 1.75, content_w, 1.05)
    text(s, MARGIN + 0.3, 1.9, content_w - 0.6, 0.8,
         ["“BM25 does not consider the semantic meaning of the query terms or the documents.”",
          [("the course, deck 5", {"color": MUTED, "size": 14, "bold": False})]],
         size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    card(s, MARGIN, 3.1, 3.8, 1.7, "345,418", "MiniLM-L6-v2 vectors over the ≥ 20-word reviews, indexed beside BM25 in Elasticsearch")
    card(s, MARGIN + 4.05, 3.1, 3.8, 1.7, "0.96", "ANN recall@10 against exact cosine on the frozen query set")
    card(s, MARGIN + 8.1, 3.1, 3.83, 1.7, "H-E1 not held", "kNN alone does not beat BM25 — client-side RRF hybrid is where the gain is", figure_size=24, accent=True)
    text(s, MARGIN, 5.15, content_w, 1.5,
         [[("Hybrid: ", {"bold": True}), ("two rankings fused client-side by reciprocal rank, Σ 1/(60 + rank).", {})],
          [("Themes: ", {"bold": True}), ("a frozen ten-theme complaint taxonomy, labelled per review by a local qwen3:8b.", {})],
          [("RAG: ", {"bold": True}), ("answers with claim-level citations over the same retriever; 30 frozen questions.", {})]],
         size=18, gap=8)

    # -- live demo divider ---------------------------------------------------------------
    s = base(prs, None, None, 280,
             "Switch to the notebook. Ten moves, one kernel: 260 s of moves + 20 s for the two "
             "surface switches. Backup if anything fails: docs/demo/ has the executed export.")
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    text(s, MARGIN, 0.7, content_w, 0.9, "Live demo", size=44, bold=True, color=WHITE)
    text(s, MARGIN, 1.55, content_w, 0.4, "Ten moves, one kernel · 4:40", size=18, color=RGBColor(0xC9, 0xCE, 0xD8))
    moves = [
        "The stack is real", "Paced replay + streaming projection", "Silver’s reconciliation identity",
        "Every run this data went through", "Decline candidates (Kibana)", "Hybrid decomposition",
        "Complaint themes for the top candidate", "The evaluation table", "RAG answers, claim by claim",
        "The stream’s own gate",
    ]
    for i, move in enumerate(moves):
        x = MARGIN + (i // 5) * 6.1
        y = 2.5 + (i % 5) * 0.85
        text(s, x, y, 0.7, 0.5, f"{i + 1:02d}", size=22, bold=True, color=ACCENT)
        text(s, x + 0.8, y + 0.05, 5.2, 0.5, move, size=18, color=WHITE)

    # 6 -- how we avoided fooling ourselves ----------------------------------------------
    s = base(prs, 6, "How we avoided fooling ourselves", 40,
             "ADR-0001: holdout + protocol freeze. ADR-0011: reproducibility gates block, quality "
             "is reported and never reopens a phase. The audit set was opened once. Two gates "
             "re-derive the tables in pandas sharing no code with the Spark jobs.")
    points = [
        ("Temporal holdout, protocol freeze", "The decline rule was developed on pre-2020 data and sealed in one commit before a single 2020–2023 point was evaluated."),
        ("Reproducibility is separate from outcome", "12 gates PASS on re-derivation; two re-derive silver and gold in pandas with no shared code (13 and 11 checks, 0 mismatched fields)."),
        ("Bars set before the measurement, never moved", "When the labeller dropped from a hosted model to a local 8B model, 0.70 stayed 0.70. The audit set was opened once."),
        ("Every number joins back to a run", "A run_id in every Iceberg snapshot, ES doc and eval artefact; the lineage gate walks 101 links."),
    ]
    for i, (head, body) in enumerate(points):
        y = 1.75 + i * 1.28
        text(s, MARGIN, y, 0.8, 0.8, f"{i + 1}", size=36, bold=True, color=ACCENT)
        text(s, MARGIN + 0.85, y + 0.02, content_w - 0.85, 0.4, head, size=20, bold=True)
        text(s, MARGIN + 0.85, y + 0.47, content_w - 0.85, 0.7, body, size=16, color=MUTED)

    # 7 -- results -----------------------------------------------------------------------
    s = base(prs, 7, "Results, zoomed out", 35,
             "The alert rate is the number; the detection rate is the caveat. Both FAILs are "
             "measured misses against bars set beforehand, published rather than hidden. "
             "'The rule finds declining products' does not follow and is not claimed.")
    card(s, MARGIN, 1.75, 5.85, 2.05, "139 of 502", "eligible products alerted in the 2020–2023 holdout (27.7%) — 148 alerts, 148 episodes", accent=True)
    card(s, MARGIN + 6.08, 1.75, 5.85, 2.05, "425 episodes", "over the whole spine — closed by recovery 207, gap 165, end of data 53")
    card(s, MARGIN, 4.0, 5.85, 2.05, "0.186 vs 0.80", "detection power at the 0.3★ step [0.168, 0.205] against the pre-registered bar — a published FAIL")
    card(s, MARGIN + 6.08, 4.0, 5.85, 2.05, "0.4583 vs 0.70", "theme macro-F1 [0.3647, 0.5223] — a published FAIL, but above the star-only floor 0.2968; MLlib baseline 0.3899")
    text(s, MARGIN, 6.3, content_w, 0.6,
         "An alert is unlikely to be noise. The 139 are a floor, not a census — and an alert is not causal proof.",
         size=18, bold=True)

    # 8 -- trade-offs and declines -------------------------------------------------------
    s = base(prs, 8, "Trade-offs, and what we declined", 30,
             "Say the four declines out loud in these words. Kafka Connect was the schedule, not "
             "the merits. Oozie is an opinion and I own it as one.")
    col = (content_w - 0.4) / 2
    text(s, MARGIN, 1.7, col, 0.3, "THREE TRADE-OFFS", size=11, bold=True, color=ACCENT)
    text(s, MARGIN, 2.1, col, 4.5,
         [[("Iceberg is not a course technology. ", {"bold": True}), ("One mention in ~860 slides; conceded up front, earned on mechanism.", {"color": MUTED})],
          [("A local qwen3:8b instead of a hosted model. ", {"bold": True}), ("The 0.70 bar deliberately unmoved, so the constraint shows as a reported miss.", {"color": MUTED})],
          [("Public but sensitive review text. ", {"bold": True}), ("Only title and text cross the wire to the model, never user_ids.", {"color": MUTED})]],
         size=18, gap=14)
    x2 = MARGIN + col + 0.4
    text(s, x2, 1.7, col, 0.3, "FOUR DECLINES, SAID OUT LOUD", size=11, bold=True, color=ACCENT)
    text(s, x2, 2.1, col, 4.8,
         [[("Kafka Connect Elasticsearch sink — dropped for time. ", {"bold": True}), ("The schedule, not the merits.", {"color": MUTED})],
          [("Single-node HDFS — declined on resources, and MinIO is not equivalent. ", {"bold": True}), ("No block replication, no atomic rename; hence Postgres holds the Iceberg catalogue.", {"color": MUTED})],
          [("Sqoop and Pig — on the course’s own slide. ", {"bold": True}), ("“Apache Sqoop moved into the Attic in June 2021”; Spark SQL does Pig’s transform.", {"color": MUTED})],
          [("Oozie — an opinion, and I own it as one. ", {"bold": True}), ("Spark and the run ledger fill the role, but Oozie is not retired.", {"color": MUTED})]],
         size=18, gap=14)

    return prs


def main() -> int:
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(TARGET)
    print(f"SLIDES_DECK slides={len(prs.slides)} path={TARGET.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
